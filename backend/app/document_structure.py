"""Text extraction that also remembers *where* each span of text came from.

The intake agent only needs a document's text, but a human checking an
extracted value needs to know which section and page it was read from. This
module produces the same text the old extraction.extract_text did -- byte for
byte, so retrieval behaves exactly as before -- alongside a list of segments
that map any character offset in that text back to a (page, section).

How reliable the two location signals are depends on the format:

- PDF: page numbers are exact. Section headings are inferred from the text
  (numbered or ALL-CAPS lines), because a PDF carries no heading structure.
- DOCX: headings come from Word heading styles when present, else from
  bold/numbered/ALL-CAPS paragraphs. Word doesn't store page numbers in the
  file, only the page breaks it last laid out, so pages here are approximate
  and only available if the file was saved by Word. `page_source` records
  which case applies so the UI never presents a guess as exact.
- TXT: no pages; headings inferred from the text.
"""
from __future__ import annotations

import bisect
import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

SECTION_SEP = " › "

PAGE_SOURCE_PDF = "pdf"
PAGE_SOURCE_DOCX_MARKERS = "docx_markers"
PAGE_SOURCE_NONE = "none"


@dataclass
class Segment:
    start: int
    end: int
    page: int | None
    section: str


@dataclass
class ExtractedDocument:
    text: str
    segments: list[Segment] = field(default_factory=list)
    page_count: int | None = None
    page_source: str = PAGE_SOURCE_NONE

    def __post_init__(self) -> None:
        self._starts = [s.start for s in self.segments]

    def locate(self, offset: int) -> tuple[int | None, str]:
        """(page, section) of the segment containing `offset`. An offset that
        falls on the newline joining two segments resolves to the earlier one."""
        if not self.segments:
            return None, ""
        i = bisect.bisect_right(self._starts, offset) - 1
        seg = self.segments[max(i, 0)]
        return seg.page, seg.section


# ---- heading detection --------------------------------------------------------------

_NUMBERED_RE = re.compile(r"^(\d+(?:\.\d+)*)[.)]\s+\S")
# "Schedule A-1", "Article IV", "Section 2.3" -- but not "Schedule Description":
# the keyword is case-insensitive, the identifier after it is not.
_KEYWORD_RE = re.compile(
    r"^(?i:section|article|schedule|exhibit|appendix|part)\s+(?:\d+(?:\.\d+)*|[A-Z]{1,2}-?\d+|[A-Z]|[IVXLC]+)\b"
)
_MARKDOWN_RE = re.compile(r"^(#{1,6})\s+(.+)$")
# Things that mean "this is data or a template blank, not a heading": brackets,
# currency/percent signs, underscore blanks, long digit runs.
_NOT_A_HEADING_RE = re.compile(r"[\[\]$%]|_{2,}|\d{5,}")


def classify_heading(line: str, *, bold: bool = False, style: str | None = None) -> tuple[int, str] | None:
    """Returns (level, title) if `line` reads as a section heading, else None.

    Deliberately conservative: a wrongly-detected heading mislabels every
    chunk beneath it, while a missed one only makes a section a little coarser.
    "Label: value" lines, sentences, placeholders and amounts are never headings.
    """
    text = " ".join(line.split())  # also folds non-breaking spaces
    if not text:
        return None

    if style:
        s = style.strip().lower()
        if s == "title":
            return 1, text
        m = re.match(r"heading\s*(\d)", s)
        if m:
            return int(m.group(1)), text

    md = _MARKDOWN_RE.match(text)
    if md:
        return len(md.group(1)), md.group(2).strip()

    if len(text) > 90 or len(text.split()) > 12:
        return None
    if text.endswith((":", ".", ";", ",")) or ":" in text:
        return None
    if _NOT_A_HEADING_RE.search(text):
        return None

    numbered = _NUMBERED_RE.match(text)
    if numbered:
        return numbered.group(1).count(".") + 1, text
    if _KEYWORD_RE.match(text):
        return 1, text

    letters = [c for c in text if c.isalpha()]
    if len(letters) >= 4 and text == text.upper() and not any(c.isdigit() for c in text):
        return 1, text
    if bold:
        return 2, text
    return None


class _SectionTracker:
    """Keeps the current heading trail (level-ordered), e.g. ["1. ELIGIBLE AR", "A/R Advance"]."""

    def __init__(self) -> None:
        self._stack: list[tuple[int, str]] = []

    def observe(self, heading: tuple[int, str] | None) -> None:
        if heading is None:
            return
        level, title = heading
        while self._stack and self._stack[-1][0] >= level:
            self._stack.pop()
        self._stack.append((level, title))

    @property
    def path(self) -> str:
        return SECTION_SEP.join(title for _, title in self._stack)


class _Builder:
    """Accumulates text parts (joined with "\\n", exactly as the old extractor
    did) and the segments describing them."""

    def __init__(self) -> None:
        self.parts: list[str] = []
        self.segments: list[Segment] = []
        self._pos = 0

    def _open_part(self, text: str) -> int:
        if self.parts:
            self._pos += 1  # the "\n" that joins this part to the previous one
        start = self._pos
        self.parts.append(text)
        self._pos += len(text)
        return start

    def add(self, text: str, page: int | None, section: str) -> None:
        start = self._open_part(text)
        self.segments.append(Segment(start, start + len(text), page, section))

    def add_lines(self, block: str, page: int | None, tracker: _SectionTracker) -> None:
        """One part, but a segment per line so headings inside it take effect
        from the line they appear on."""
        block_start = self._open_part(block)
        offset = 0
        for raw in block.split("\n"):
            tracker.observe(classify_heading(raw))
            self.segments.append(Segment(block_start + offset, block_start + offset + len(raw), page, tracker.path))
            offset += len(raw) + 1

    def build(self, page_count: int | None, page_source: str) -> ExtractedDocument:
        return ExtractedDocument("\n".join(self.parts), self.segments, page_count, page_source)


# ---- per-format extraction ------------------------------------------------------------

def extract_document(filename: str, content: bytes) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(content)
    if suffix == ".docx":
        return _extract_docx(content)
    return from_plain_text(content.decode("utf-8", errors="ignore"))


def from_plain_text(text: str) -> ExtractedDocument:
    builder = _Builder()
    builder.add_lines(text, None, _SectionTracker())
    return builder.build(None, PAGE_SOURCE_NONE)


def _extract_pdf(content: bytes) -> ExtractedDocument:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    builder = _Builder()
    tracker = _SectionTracker()
    for number, page in enumerate(reader.pages, start=1):
        builder.add_lines(page.extract_text() or "", number, tracker)
    return builder.build(len(reader.pages), PAGE_SOURCE_PDF)


def _is_bold(paragraph) -> bool:
    runs = [r for r in paragraph.runs if r.text.strip()]
    if not runs:
        return False
    style_bold = bool(paragraph.style is not None and paragraph.style.font is not None and paragraph.style.font.bold)
    return all((r.bold if r.bold is not None else style_bold) for r in runs)


def _extract_docx(content: bytes) -> ExtractedDocument:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(BytesIO(content))
    body = document.element.body

    # Word writes <w:lastRenderedPageBreak/> wherever it last started a new page.
    # An explicit page break usually gets one too, so counting both would double
    # count -- use the rendered markers when the file has any, else explicit breaks.
    has_rendered = bool(body.xpath(".//w:lastRenderedPageBreak"))
    has_explicit = bool(body.xpath('.//w:br[@w:type="page"]'))
    paginated = has_rendered or has_explicit
    break_tag = qn("w:lastRenderedPageBreak") if has_rendered else qn("w:br")
    text_tag = qn("w:t")
    type_attr = qn("w:type")

    page = 1
    tracker = _SectionTracker()

    def scan(element) -> int | None:
        """Advances `page` across any breaks in `element` (document order) and
        returns the page its first text sits on."""
        nonlocal page
        first_text_page: int | None = None
        for node in element.iter():
            if node.tag == break_tag and (has_rendered or node.get(type_attr) == "page"):
                page += 1
            elif node.tag == text_tag and node.text and node.text.strip() and first_text_page is None:
                first_text_page = page
        return first_text_page if first_text_page is not None else page

    paragraph_texts: list[tuple[str, int | None, str]] = []
    row_texts: list[tuple[str, int | None, str]] = []

    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            paragraph_page = scan(child)
            style_name = paragraph.style.name if paragraph.style is not None else None
            if "\n" not in paragraph.text:
                tracker.observe(classify_heading(paragraph.text, bold=_is_bold(paragraph), style=style_name))
            paragraph_texts.append((paragraph.text, paragraph_page if paginated else None, tracker.path))
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row_element, row in zip(child.iterchildren(qn("w:tr")), table.rows):
                row_page = scan(row_element)
                row_texts.append((" | ".join(cell.text for cell in row.cells), row_page if paginated else None, tracker.path))

    # Emit in the extractor's long-standing order (all paragraphs, then all table
    # rows) so the text -- and therefore retrieval -- is unchanged; each part
    # keeps the page/section of where it really sits in the document.
    builder = _Builder()
    for text, part_page, section in paragraph_texts + row_texts:
        builder.add(text, part_page, section)
    return builder.build(page if paginated else None, PAGE_SOURCE_DOCX_MARKERS if paginated else PAGE_SOURCE_NONE)
